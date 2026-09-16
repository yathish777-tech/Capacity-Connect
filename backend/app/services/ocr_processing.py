"""
Extracts plain text from course materials so it can be chunked and embedded.

- PDFs: PyMuPDF (fitz) text layer first; falls back to Tesseract OCR per page
  when a page has no extractable text (i.e. it's a scanned image).
- PPTX: python-pptx, reading every text frame on every slide.
"""

import io

import fitz  # PyMuPDF
import pytesseract
from PIL import Image
from pptx import Presentation


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts: list[str] = []
    with fitz.open(stream=file_bytes, filetype="pdf") as doc:
        for page in doc:
            page_text = page.get_text().strip()
            if page_text:
                text_parts.append(page_text)
            else:
                # Scanned page with no text layer -> OCR it.
                pix = page.get_pixmap(dpi=200)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                text_parts.append(pytesseract.image_to_string(image))
    return "\n\n".join(text_parts)


def extract_text_from_pptx(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    text_parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                slide_text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text)
                if slide_text:
                    text_parts.append(slide_text)
    return "\n\n".join(text_parts)


def extract_text(file_bytes: bytes, file_type: str) -> str:
    if file_type == "pdf":
        return extract_text_from_pdf(file_bytes)
    if file_type == "ppt":
        return extract_text_from_pptx(file_bytes)
    raise ValueError(f"No text extractor for file_type={file_type!r} (only pdf/ppt are auto-ingested for RAG)")


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 150) -> list[str]:
    """Simple sliding-window chunker on characters, breaking on paragraph boundaries where possible."""
    text = text.strip()
    if not text:
        return []
    overlap = min(overlap, chunk_size // 2)

    chunks: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + chunk_size, n)
        if end < n:
            boundary = text.rfind("\n\n", start, end)
            if boundary != -1 and boundary > start + chunk_size // 2:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = end - overlap
    return chunks
